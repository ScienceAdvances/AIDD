"""Read original CAR-T workbooks without modifying them; export cell-traceable observations."""
from pathlib import Path
import argparse
import csv
import json
import math
import statistics
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data/cart"
OUT = ROOT / "examples/debian13_20261023/cart"


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data", type=Path, default=DATA)
    parser.add_argument("--output", type=Path, default=OUT)
    args = parser.parse_args()
    data, out = args.data.resolve(), args.output.resolve()
    out.mkdir(parents=True, exist_ok=True)
    names = ["Source data - Main figures.xlsx", "Source data - Supplementary figures.xlsx"]
    books = [openpyxl.load_workbook(data / name, read_only=True, data_only=True) for name in names]
    rows = []
    panels = [
        (0, "Figure 1", "1E", "A", "B", 2, 31, "relative_target_viability_pct"),
        (0, "Figure 1", "1F", "D", "E", 2, 31, "relative_target_viability_pct"),
        (0, "Figure 2", "2B", "A", "B", 2, 19, "CAR_surface_expression_MFI"),
        (0, "Figure 2", "2I", "G", "H", 2, 25, "IFNg_source_value_unit_not_in_sheet"),
        (0, "Figure 4", "4C", "A", "B", 2, 29, "CAR_positive_in_transduced_cells_pct"),
        (0, "Figure 4", "4H", "J", "K", 2, 25, "relative_IFNg_level"),
        (1, "Fig. S7", "S7B", "D", "E", 2, 15, "antigen_binding_MFI"),
    ]
    for book, sheet, panel, label_col, value_col, first, last, metric in panels:
        s = books[book][sheet]
        group = None
        group_cell = None
        observation = 0
        for r in range(first, last + 1):
            label = s[f"{label_col}{r}"].value
            value = s[f"{value_col}{r}"].value
            if label is not None:
                group, group_cell, observation = str(label), f"{label_col}{r}", 0
            if isinstance(value, (int, float)):
                observation += 1
                rows.append({"panel": panel, "metric": metric, "group": group, "observation_index": observation,
                             "value": value, "source_book": names[book], "source_sheet": sheet,
                             "source_cell": f"{value_col}{r}", "group_label_cell": group_cell,
                             "donor_id": "not_provided"})
    s = books[1]["Fig. S7"]
    for col, group in [(5, "CAR[B]"), (8, "[B]-H2"), (11, "[B]-L3"), (14, "[B]-H2L3wt"), (17, "CAR[E]"), (20, "Mock")]:
        for row in [23, 24]:
            for i in range(3):
                c = s.cell(row, col + i)
                rows.append({"panel": f"S7I_{s.cell(row,4).value}", "metric": "relative_target_viability_pct", "group": group,
                             "observation_index": i + 1, "value": c.value, "source_book": names[1], "source_sheet": "Fig. S7",
                             "source_cell": c.coordinate, "group_label_cell": s.cell(22, col).coordinate, "donor_id": "not_provided"})
    assert all(row["group"] and isinstance(row["value"], (int, float)) for row in rows)
    with (out / "source_observations.csv").open("w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0]))
        writer.writeheader(); writer.writerows(rows)
    groups = {}
    for row in rows:
        groups.setdefault((row["panel"], row["metric"], row["group"]), []).append(row["value"])
    summaries = []
    for (panel, metric, group), values in groups.items():
        n = len(values)
        sd = statistics.stdev(values) if n > 1 else None
        summaries.append({"panel": panel, "metric": metric, "group": group, "n_source_observations": n,
                          "mean": statistics.mean(values), "sd_sample": sd, "sem": sd / math.sqrt(n) if sd is not None else None,
                          "minimum": min(values), "maximum": max(values)})
    with (out / "panel_summary.csv").open("w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(summaries[0]))
        writer.writeheader(); writer.writerows(summaries)
    inventory = {name: [{"sheet": sh.title, "rows": sh.max_row, "columns": sh.max_column} for sh in book]
                 for name, book in zip(names, books)}
    summary = {"observations": len(rows), "panel_groups": len(summaries), "source_books_modified": False,
               "new_hypothesis_tests": False, "donor_pairing_assumed": False, "input_directory": str(data), "sheets": inventory}
    (out / "extraction_summary.json").write_text(json.dumps(summary, indent=2) + "\n")
    for r in summaries:
        if r["panel"] in ("1F", "4H", "S7I_5:1"):
            print(r["panel"], r["group"], r["n_source_observations"], round(r["mean"], 4), round(r["sem"], 4))


if __name__ == "__main__":
    main()
